#!/usr/bin/env python3
"""Build the submission deck from the official exported Databricks template."""

from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT


SLIDES = (
    (14, (
        "From Fragmented Signals to Timely Student Support",
        "A governed student-retention workflow on Databricks",
        "Public Sector • Higher Education • Synthetic demonstration",
    )),
    (29, (
        "Advisors need to know who needs attention today—and why—without compromising trust.",
    )),
    (39, (
        "Buyer KPIs and synthetic demo baseline",
        "Designed values demonstrate the workflow—not real institutional performance.",
        "88.88%\nSynthetic retention rate",
        "2,664\nInitially elevated-risk students",
        "$24.1M\nEstimated synthetic tuition exposure",
        "20,000\nCurrent synthetic students",
    )),
    (34, (
        "Current-state friction delays action",
        "The problem is operational, not simply predictive.",
        "Fragmented signals",
        "Inconsistent prioritization",
        "Weak feedback loop",
        "Attendance, LMS, finance, enrollment, and case data arrive on different clocks.",
        "Advisors reconcile data manually and lack a defensible daily queue.",
        "Outreach and outcomes are disconnected from executive coverage and analytics.",
    )),
    (23, (
        "A closed-loop workflow from signal to action",
        "One governed path connects observation, human action, and analytical learning.",
        "1. Observe — trigger one deterministic synthetic day\n"
        "2. Govern — apply quality, lineage, grants, and row filters\n"
        "3. Prioritize — score current snapshots with a versioned ML model\n"
        "4. Assist — generate cited, evaluated advisor briefings\n"
        "5. Act — record interventions transactionally in Lakebase\n"
        "6. Learn — return CDC to governed Gold\n"
        "7. Understand — ask embedded Genie and inspect its SQL",
    )),
    (25, (
        "Business value for two accountable owners",
        "A single experience serves different decisions without collapsing access boundaries.",
        "Executive sponsor",
        "Student-success owner",
        "See aggregate risk concentration, intervention coverage, retention, and estimated tuition exposure—without student detail.",
        "Reduce signal-to-outreach time with a prioritized caseload, grounded context, and closed-loop follow-up history.",
    )),
    (24, (
        "Governed Databricks architecture",
        "Serving refreshes cannot overwrite advisor transactions.",
        "Synthetic signals → Lakeflow Bronze / Silver → MLflow scoring → "
        "Unity Catalog Gold → read-only Lakebase serving → Databricks App + "
        "embedded Genie → Lakebase-owned interventions → CDC → governed "
        "intervention Gold",
    )),
    (34, (
        "Hybrid intelligence keeps each decision explainable",
        "Use each capability for what it does best.",
        "ML classifier",
        "Grounded GenAI",
        "Human advisor",
        "Repeatable prioritization, calibrated scores, stable tiers, version history, and baseline comparison.",
        "Allowlisted facts, citations, bounded discussion prompts, and fixed-set evaluation.",
        "Judgment, outreach, reconciliation, and accountable action remain with people.",
    )),
    (23, (
        "Live demo: one coherent student-success story",
        "Expected duration: 12–15 minutes; no manual data repair.",
        "Replay the September 23 synthetic incident day\n"
        "Inspect quality and feature freshness\n"
        "Show a medium-to-high risk transition\n"
        "Review the governed caseload and cited briefing\n"
        "Record outreach and a follow-up\n"
        "Verify Lakebase CDC and intervention Gold\n"
        "View aggregate executive coverage\n"
        "Ask embedded Genie and inspect generated SQL",
    )),
    (39, (
        "Execution evidence—not screenshots",
        "Every build domain has committed, text-readable proof.",
        "60,000\nImmutable daily risk-history rows",
        "0.9669\nSynthetic validation PR-AUC",
        "SUCCESS\nWorkflow and idempotent replay",
        "ONLINE\nBoth serving sync tables",
    )),
    (25, (
        "Trust is designed in—and limitations stay visible",
        "Responsible use means exposing both controls and boundaries.",
        "Trust controls",
        "Honest boundaries",
        "Row-filtered advisor access\nAggregate-only executive products\nRead-only serving tables\nCited AI and inspectable SQL\nVisible stale, partial, and error states",
        "Synthetic data and designed signals\nNo causal impact claim\nNo automated adverse action\nNo institutional validation yet\nTriggered daily—not real time",
    )),
    (15, (
        "Pilot the workflow—not just the model",
        "Align owners • validate institutional data and thresholds • measure signal-to-outreach time and intervention coverage",
        "Governed decisions. Accountable action. Measurable learning.",
    )),
)


def _clone_slide(presentation: Presentation, source):
    destination = presentation.slides.add_slide(presentation.slide_layouts[10])
    for shape in tuple(destination.shapes):
        destination.shapes._spTree.remove(shape.element)

    relationship_map: dict[str, str] = {}
    for relationship in source.part.rels.values():
        if relationship.reltype in {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE}:
            continue
        if relationship.is_external:
            new_id = destination.part.rels.get_or_add_ext_rel(
                relationship.reltype, relationship.target_ref
            )
        else:
            new_id = destination.part.rels.get_or_add(
                relationship.reltype, relationship.target_part
            )
        relationship_map[relationship.rId] = new_id

    for shape in source.shapes:
        element = deepcopy(shape.element)
        for child in element.iter():
            for attribute, value in tuple(child.attrib.items()):
                if value in relationship_map:
                    child.attrib[attribute] = relationship_map[value]
        destination.shapes._spTree.insert_element_before(element, "p:extLst")

    source_background = source._element.cSld.bg
    if source_background is not None:
        destination._element.cSld.insert(0, deepcopy(source_background))
    return destination


def _text_shapes(slide):
    return sorted(
        (
            shape
            for shape in slide.shapes
            if hasattr(shape, "text_frame") and shape.text.strip() and shape.text != "‹#›"
        ),
        key=lambda shape: (shape.top, shape.left),
    )


def _set_text(slide, values: tuple[str, ...], slide_number: int) -> None:
    shapes = _text_shapes(slide)
    if len(shapes) != len(values):
        raise ValueError(
            f"Slide {slide_number} expects {len(values)} text shapes, found {len(shapes)}"
        )
    for shape, value in zip(shapes, values, strict=True):
        shape.text = value
    for shape in slide.shapes:
        if hasattr(shape, "text_frame") and shape.text == "‹#›":
            shape.text = str(slide_number)


def _remove_original_slides(presentation: Presentation, count: int) -> None:
    slide_ids = presentation.slides._sldIdLst
    for slide_id in list(slide_ids)[:count]:
        presentation.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


def build(template: Path, output: Path) -> None:
    presentation = Presentation(template)
    original_count = len(presentation.slides)
    original_slides = tuple(presentation.slides)
    for number, (source_index, values) in enumerate(SLIDES, start=1):
        slide = _clone_slide(presentation, original_slides[source_index])
        _set_text(slide, values, number)
    _remove_original_slides(presentation, original_count)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)


def validate(path: Path) -> None:
    presentation = Presentation(path)
    if len(presentation.slides) != len(SLIDES):
        raise ValueError(f"Expected {len(SLIDES)} slides, found {len(presentation.slides)}")
    forbidden = (
        "Corporate Slide Deck Template",
        "Your subtitle here",
        "Your content goes here",
        "This slide is an all-purpose slide",
        "These cards are",
        "Here’s a column subheader",
        "‹#›",
    )
    all_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text_frame")
    )
    remaining = [value for value in forbidden if value in all_text]
    if remaining:
        raise ValueError(f"Template placeholders remain: {remaining}")
    if "synthetic" not in all_text.lower():
        raise ValueError("Synthetic-data disclosure is missing")


def _arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


if __name__ == "__main__":
    args = _arguments()
    build(args.template, args.output)
    validate(args.output)
    print(args.output.resolve())
